#!/usr/bin/env python3
"""Solve DBA5103 Assignment 1 and generate a polished DOCX submission.

The script is intentionally self-contained. It reads the supplied portfolio data,
solves the numerical linear program, creates every equation and chart with
Matplotlib, builds the Word document, and removes temporary image files.
"""

from __future__ import annotations

import csv
import shutil
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from scipy.optimize import linprog

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "sources" / "investment.csv"
NUS_LOGO_FILE = ROOT / "sources" / "nus_logo.png"
OUTPUT_FILE = ROOT / "DBA5103_Assignment_1_Solutions.docx"

NAVY = "18354A"
BLUE = "2F6483"
PALE_GREY = "F3F5F6"
MID_GREY = "66717A"
WHITE = "FFFFFF"
BLACK = "111820"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep and keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)
    elif not keep and keep_next is not None:
        p_pr.remove(keep_next)


def set_repeat_on_each_page(section, text: str) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Header"
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Footer"
    run = p.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def add_bottom_border(paragraph, color=NAVY, size="10", space="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)
    # set_repeat_on_each_page(section, "DBA5103  |  Operations Research and Analytics  |  Assignment 1")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(NAVY)

    styles["Title"].font.size = Pt(28)
    styles["Title"].font.bold = True
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(MID_GREY)

    h1 = styles["Heading 1"]
    h1.font.size = Pt(17)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(7)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    if "Formula Caption" not in styles:
        style = styles.add_style("Formula Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(8)
        style.font.italic = True
        style.font.color.rgb = RGBColor.from_string(MID_GREY)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_after = Pt(5)

    if "Code" not in styles:
        style = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Courier New"
        style.font.size = Pt(7.8)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_text(doc: Document, text: str, *, bold_lead: str | None = None, italic=False) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        r = p.add_run(text)
        r.italic = italic


def add_answer_box(doc: Document, label: str, text: str) -> None:
    """Add a labelled answer as ordinary body text, without a shaded box."""
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    p.add_run(text)


def render_formula(path: Path, lines: str | list[str], fontsize: float = 16.5) -> None:
    if isinstance(lines, str):
        lines = [lines]
    height = max(0.44, 0.42 * len(lines))
    fig = plt.figure(figsize=(7.1, height), dpi=300, facecolor="white")
    for idx, line in enumerate(lines):
        y = 1 - (idx + 0.5) / len(lines)
        fig.text(0.5, y, line, ha="center", va="center", fontsize=fontsize, color="#111820")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.07, facecolor="white")
    plt.close(fig)


def render_matrix(path: Path, values, prefix="A =", fontsize=17.0) -> None:
    rows = len(values)
    cols = len(values[0])
    fig_w = 1.8 + cols * 0.48
    fig_h = 0.35 + rows * 0.45
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=300)
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)
    ax.axis("off")
    y_mid = fig_h / 2
    prefix_x = 0.55
    left = 1.15
    right = left + cols * 0.48 + 0.12
    top = fig_h - 0.12
    bottom = 0.12
    ax.text(prefix_x, y_mid, prefix, ha="center", va="center", fontsize=fontsize)
    ax.plot([left, left, left + 0.12], [top, bottom, bottom], color=f"#{BLACK}", lw=1.5)
    ax.plot([left, left + 0.12], [top, top], color=f"#{BLACK}", lw=1.5)
    ax.plot([right, right, right - 0.12], [top, bottom, bottom], color=f"#{BLACK}", lw=1.5)
    ax.plot([right, right - 0.12], [top, top], color=f"#{BLACK}", lw=1.5)
    for i, row in enumerate(values):
        y = top - (i + 0.5) * (top - bottom) / rows
        for j, val in enumerate(row):
            x = left + 0.25 + j * 0.48
            ax.text(x, y, str(val), ha="center", va="center", fontsize=fontsize)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.06, facecolor="white")
    plt.close(fig)


def add_formula(doc: Document, image_path: Path, caption: str | None = None, max_width=6.15) -> None:
    with Image.open(image_path) as im:
        dpi = im.info.get("dpi", (300, 300))[0] or 300
        natural_width = im.width / dpi
    width = min(max_width, max(1.15, natural_width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2 if caption else 5)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph(caption, style="Formula Caption")
        cp.paragraph_format.keep_with_next = False


def add_figure(doc: Document, image_path: Path, caption: str, width=4.65) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cp = doc.add_paragraph(caption, style="Formula Caption")
    cp.paragraph_format.keep_with_next = False


def make_table(doc: Document, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = widths is None
    header = table.rows[0]
    set_repeat_table_header(header)
    for j, heading in enumerate(headers):
        cell = header.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(heading))
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Inches(widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_margins(cell)
            if i % 2 == 1:
                set_cell_shading(cell, PALE_GREY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.name = "Arial"
            r.font.size = Pt(font_size)
            if widths:
                cell.width = Inches(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def plot_triangle(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(4.7, 3.7), dpi=240)
    triangle = np.array([[0, 0], [1, 0], [0, 1]])
    ax.fill(triangle[:, 0], triangle[:, 1], color="#BFD8E6", alpha=0.9, label="Feasible set")
    ax.plot([0, 1], [1, 0], color="#18354A", linewidth=2.1, label="x + y = 1")
    ax.scatter([0, 1, 0], [0, 0, 1], color="#18354A", s=22, zorder=3)
    ax.annotate("(1, 0)", (1, 0), xytext=(-5, 9), textcoords="offset points", ha="right", fontsize=9)
    ax.annotate("(0, 1)", (0, 1), xytext=(7, -2), textcoords="offset points", fontsize=9)
    ax.set(xlim=(0, 1.14), ylim=(0, 1.14), xlabel="x", ylabel="y")
    ax.set_aspect("equal")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(alpha=0.18)
    ax.legend(frameon=False, loc="upper right", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_quarter_disk(path: Path) -> None:
    theta = np.linspace(0, np.pi / 2, 300)
    x, y = np.cos(theta), np.sin(theta)
    fig, ax = plt.subplots(figsize=(4.7, 3.7), dpi=240)
    px = np.r_[0, x, 0]
    py = np.r_[0, y, 0]
    ax.fill(px, py, color="#D8E6BE", alpha=0.95, label="Feasible set")
    ax.plot(x, y, color="#3F6835", linewidth=2.1, label="x² + y² = 1")
    ax.scatter([0, 1, 0], [0, 0, 1], color="#3F6835", s=22, zorder=3)
    ax.set(xlim=(0, 1.14), ylim=(0, 1.14), xlabel="x", ylabel="y")
    ax.set_aspect("equal")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(alpha=0.18)
    ax.legend(frameon=False, loc="upper right", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def read_portfolio(path: Path):
    records = []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            records.append(
                {
                    "stock": row["Stock"],
                    "p": float(row["Buying Price"]),
                    "s": float(row["Current Share"]),
                    "q": float(row["Current Price"]),
                    "r": float(row["Expected Future Price"]),
                }
            )
    return records


def solve_investment(records, target=9000.0):
    p = np.array([row["p"] for row in records], dtype=float)
    s = np.array([row["s"] for row in records], dtype=float)
    q = np.array([row["q"] for row in records], dtype=float)
    r = np.array([row["r"] for row in records], dtype=float)
    net = 0.99 * q - 0.30 * np.maximum(q - p, 0.0)
    result = linprog(
        c=r,
        A_ub=-net.reshape(1, -1),
        b_ub=np.array([-target]),
        bounds=[(0.0, float(limit)) for limit in s],
        method="highs",
    )
    if not result.success:
        raise RuntimeError(f"Investment LP failed: {result.message}")
    sold = result.x
    cash = net * sold
    remaining = s - sold
    future_by_stock = r * remaining
    future_total = float(np.sum(future_by_stock))
    return {
        "p": p,
        "s": s,
        "q": q,
        "r": r,
        "net": net,
        "ratio": r / net,
        "sold": sold,
        "cash": cash,
        "remaining": remaining,
        "future_by_stock": future_by_stock,
        "future_total": future_total,
        "cash_total": float(np.sum(cash)),
        "solver_message": result.message,
    }


def plot_portfolio_solution(path: Path, records, solution) -> None:
    labels = [row["stock"] for row in records]
    y = np.arange(len(labels))
    sold = solution["sold"]
    retained = solution["remaining"]
    fig, ax = plt.subplots(figsize=(6.5, 3.55), dpi=240)
    ax.barh(y, sold, color="#2F6483", label="Sold")
    ax.barh(y, retained, left=sold, color="#C8D6DE", label="Retained")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlabel("Shares")
    ax.set_xlim(0, max(solution["s"]) * 1.04)
    ax.grid(axis="x", alpha=0.2)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, loc="lower right")
    for idx, value in enumerate(sold):
        if value > 0.01:
            label = f"{value:,.0f}" if abs(value - round(value)) < 0.01 else f"{value:,.2f}"
            ax.text(min(value / 2, 500), idx, label, ha="center", va="center", color="white", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def get_nus_logo() -> Path:
    """Validate and return the locally stored complete NUS logo."""
    if not NUS_LOGO_FILE.exists():
        raise FileNotFoundError(f"NUS logo is missing: {NUS_LOGO_FILE}")
    with Image.open(NUS_LOGO_FILE) as logo:
        if logo.width < 300 or logo.height < 100:
            raise ValueError("NUS logo is unexpectedly small")
        logo.verify()
    return NUS_LOGO_FILE


def add_title_page(doc: Document) -> None:
    logo_path = get_nus_logo()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(24)
    p.add_run().add_picture(str(logo_path), width=Inches(2.95))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("DBA5103")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph("Operations Research and Analytics", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)

    p = doc.add_paragraph("Assignment 1", style="Title")
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run("Semester I, Academic Year 2026/2027")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Reagan")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("MSBA | NUS")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Student No: A0319086X")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_question_1(doc: Document, image_dir: Path) -> None:
    doc.add_heading("1. Production Planning", level=1)
    doc.add_heading("1(a)(i) Linear programming model", level=2)
    add_text(
        doc,
        "Using the production and procurement variables defined in the question, a correct linear program is:",
    )
    p = image_dir / "q1_scalar.png"
    render_formula(
        p,
        [
            r"$\max\quad \sum_{j=1}^{n}p_jx_j-\sum_{i=1}^{m}c_iy_i$",
            r"$\mathrm{s.t.}\quad \sum_{j=1}^{n}a_{ij}x_j\leq y_i\quad(i=1,\ldots,m)$",
            r"$\sum_{i=1}^{m}c_iy_i\leq B,\qquad x_j\geq0,\quad y_i\geq0.$",
        ],
    )
    add_formula(doc, p)
    add_text(
        doc,
        "The first set of constraints ensures that enough of each raw material is procured, and the second constraint enforces the budget.",
    )

    doc.add_heading("1(a)(ii) Matrix form", level=2)
    p = image_dir / "q1_matrix.png"
    render_formula(
        p,
        [
            r"$\max_{x,y}\quad p^{\mathsf{T}}x-c^{\mathsf{T}}y$",
            r"$\mathrm{s.t.}\quad Ax-y\leq0,\qquad c^{\mathsf{T}}y\leq B,\qquad x\geq0,\ y\geq0.$",
        ],
    )
    add_formula(doc, p)
    add_text(doc, "Equivalently, the variables and constraints can be stacked into one block model.")
    p = image_dir / "q1_block.png"
    render_formula(
        p,
        [
            r"$\max\quad [\ p^{\mathsf{T}}\ \ -c^{\mathsf{T}}\ ]\binom{x}{y}$",
            r"$\mathrm{s.t.}\quad \left[\ A\ \ -I_m\ \right]\binom{x}{y}\leq0,\qquad [\ 0\ \ c^{\mathsf{T}}\ ]\binom{x}{y}\leq B,$",
            r"$x\geq0,\qquad y\geq0.$",
        ],
        fontsize=15.5,
    )
    add_formula(doc, p)

    doc.add_heading("1(b) Model without procurement variables", level=2)
    add_text(doc, "Substituting material usage for procurement removes the procurement variables and gives:")
    p = image_dir / "q1_reduced.png"
    render_formula(
        p,
        [
            r"$\max_{x}\quad (p-A^{\mathsf{T}}c)^{\mathsf{T}}x$",
            r"$\mathrm{s.t.}\quad c^{\mathsf{T}}Ax\leq B,\qquad x\geq0.$",
        ],
    )
    add_formula(doc, p)
    add_answer_box(doc, "Interpretation", "The unit profit vector is the selling-price vector less the raw-material cost attributable to one unit of each product.")


def add_question_2(doc: Document, image_dir: Path) -> None:
    doc.add_heading("2. Matrix Completion", level=1)
    doc.add_heading("2(a) Rank of the given matrix", level=2)
    p = image_dir / "q2_rank_relation.png"
    render_formula(
        p,
        [
            r"$R_3=R_1+R_2,$",
            r"$R_1\ \mathrm{and}\ R_2\ \mathrm{are\ linearly\ independent}\quad\Longrightarrow\quad \operatorname{rank}(A)=2.$",
        ],
    )
    add_formula(doc, p)
    add_text(doc, "The third row therefore adds no new dimension, while the first two rows establish two independent directions.")

    doc.add_heading("2(b) Rank of an outer product", level=2)
    p = image_dir / "q2_outer.png"
    render_formula(
        p,
        [
            r"$A=xy^{\mathsf{T}},\qquad A_{:j}=y_jx.$",
            r"$\operatorname{col}(A)\subseteq\operatorname{span}\{x\},\qquad x\neq0,\ y\neq0,$",
            r"$\therefore\qquad \operatorname{rank}(A)=1.$",
        ],
    )
    add_formula(doc, p)
    add_text(doc, "Every column is a scalar multiple of the same nonzero vector. Since the second vector is nonzero, at least one column is nonzero, so the rank cannot be zero.")

    doc.add_heading("2(c) Completion of the ratings matrix", level=2)
    add_text(
        doc,
        "The observed one in the first row and first column makes both associated factors nonzero. The observed zero in column 2 therefore forces the second movie factor to zero. The observed one in column 3 equates its factor to that of column 1. The observations in the third row then make the third user factor equal to the first, force the column 4 factor to zero, and equate the column 5 factor to the column 1 factor. Finally, the zero in row 2 and column 3 forces the second user factor to zero.",
    )
    p = image_dir / "q2_completion_reason.png"
    render_formula(
        p,
        [
            r"$x_1y_1=1,\quad y_2=0,\quad y_3=y_1,\quad x_3=x_1,\quad y_4=0,\quad y_5=y_1,\quad x_2=0.$",
        ],
        fontsize=15.3,
    )
    add_formula(doc, p)
    p = image_dir / "q2_completed_matrix.png"
    render_matrix(p, [[1, 0, 1, 0, 1], [0, 0, 0, 0, 0], [1, 0, 1, 0, 1]])
    add_formula(doc, p, caption="Completed rank-one ratings matrix")
    add_answer_box(doc, "Check", "Rows 1 and 3 are identical and nonzero, while row 2 is zero. The completed matrix consequently has rank one and matches every observed entry.")


def add_question_3(doc: Document, image_dir: Path) -> None:
    doc.add_heading("3. Polyhedra", level=1)
    doc.add_heading("3(a) Parameterized linear feasible set", level=2)
    add_text(
        doc,
        "Yes. For every fixed nonnegative value of the parameter, all coefficients are constants, so the set is the intersection of three linear halfspaces. It is therefore a polyhedron. When the parameter is zero, the first inequality is redundant and the feasible set is the unbounded first quadrant, which remains a polyhedron.",
    )
    p = image_dir / "q3a_intercepts.png"
    render_formula(
        p,
        [
            r"$a>0:\qquad a^2x+\sqrt{a}\,y\leq1,\qquad 0\leq x\leq a^{-2},\quad 0\leq y\leq a^{-1/2}.$",
            r"$a=1:\qquad x+y\leq1,\qquad x\geq0,\quad y\geq0.$",
        ],
    )
    add_formula(doc, p)
    plot_path = image_dir / "q3a_triangle.png"
    plot_triangle(plot_path)
    add_figure(doc, plot_path, "Figure 1. Feasible set for a = 1", width=4.15)

    doc.add_heading("3(b) Quarter unit disk", level=2)
    add_text(
        doc,
        "The feasible region is the quarter of the closed unit disk in the first quadrant. It is not a polyhedron because its circular arc cannot be generated by a finite number of linear halfspace boundaries.",
    )
    plot_path = image_dir / "q3b_quarter_disk.png"
    plot_quarter_disk(plot_path)
    add_figure(doc, plot_path, "Figure 2. Feasible quarter unit disk", width=4.15)

    doc.add_heading("3(c) Quadratic inequality with a bounded solution set", level=2)
    p = image_dir / "q3c.png"
    render_formula(
        p,
        [r"$x^2-8x+15=(x-3)(x-5)\leq0\quad\Longleftrightarrow\quad 3\leq x\leq5.$"],
    )
    add_formula(doc, p)
    add_answer_box(doc, "Answer", "Yes. The interval is the intersection of the two halfspaces defined by its lower and upper bounds, so it is a polyhedron in one dimension.")

    doc.add_heading("3(d) Quadratic inequality with a disconnected solution set", level=2)
    p = image_dir / "q3d.png"
    render_formula(
        p,
        [r"$x^2-8x+15\geq0\quad\Longleftrightarrow\quad x\leq3\ \ \mathrm{or}\ \ x\geq5.$"],
    )
    add_formula(doc, p)
    add_answer_box(doc, "Answer", "No. The union of the two rays is not convex. Every polyhedron, as an intersection of halfspaces, must be convex.")


def add_question_4(doc: Document, image_dir: Path) -> None:
    doc.add_heading("4. Manpower Scheduling", level=1)
    add_text(
        doc,
        "Let M, A, and E denote the numbers of full-time clerks starting the morning, afternoon, and evening shifts. Let P1, P2, P3, and P4 denote part-time clerks assigned to the four consecutive four-hour periods. All variables are continuous and nonnegative, as requested.",
    )
    make_table(
        doc,
        ["Period", "Full-time clerks on duty", "Part-time variable", "Minimum staff"],
        [
            ["8 A.M. to noon", "M", "P1", "4"],
            ["Noon to 4 P.M.", "M + A", "P2", "8"],
            ["4 P.M. to 8 P.M.", "A + E", "P3", "10"],
            ["8 P.M. to midnight", "E", "P4", "6"],
        ],
        widths=[1.55, 1.7, 1.35, 1.25],
        font_size=8.3,
    )
    add_text(doc, "An eight-hour full-time shift costs $320 and a four-hour part-time shift costs $120. The objective is:")
    p = image_dir / "q4_objective.png"
    render_formula(p, r"$\min\quad 320(M+A+E)+120(P_1+P_2+P_3+P_4)$")
    add_formula(doc, p)

    add_text(doc, "The minimum coverage requirements are:")
    p = image_dir / "q4_coverage.png"
    render_formula(
        p,
        [
            r"$M+P_1\geq4,\qquad M+A+P_2\geq8,$",
            r"$A+E+P_3\geq10,\qquad E+P_4\geq6.$",
        ],
    )
    add_formula(doc, p)

    add_text(doc, "The full-time staffing requirement in each period is:")
    p = image_dir / "q4_mix.png"
    render_formula(
        p,
        [
            r"$M\geq P_1,\qquad M+A\geq P_2,$",
            r"$A+E\geq P_3,\qquad E\geq P_4.$",
        ],
    )
    add_formula(doc, p)
    p = image_dir / "q4_nonnegative.png"
    render_formula(p, r"$M,A,E,P_1,P_2,P_3,P_4\geq0.$")
    add_formula(doc, p)
    add_answer_box(doc, "LP status", "The objective, coverage constraints, staffing-mix constraints, and nonnegativity restrictions are all linear. Integer restrictions have been omitted as instructed.")


def add_question_5(doc: Document, image_dir: Path) -> None:
    doc.add_heading("5. Absolute-Value Reformulation", level=1)
    add_text(
        doc,
        "Introduce nonnegative auxiliary variables t, u, and v to represent the three absolute-value expressions. Keep the original decision variables unrestricted in sign. An equivalent linear program is:",
    )
    p = image_dir / "q5_lp.png"
    render_formula(
        p,
        [
            r"$\min\quad 2x_1+3t$",
            r"$\mathrm{s.t.}\quad t\geq x_2-10,\qquad t\geq10-x_2,$",
            r"$u\geq x_1+2,\qquad u\geq-x_1-2,$",
            r"$v\geq x_2,\qquad v\geq-x_2,\qquad u+v\leq5,$",
            r"$t,u,v\geq0,\qquad x_1,x_2\ \mathrm{unrestricted}.$",
        ],
        fontsize=15.0,
    )
    add_formula(doc, p)
    add_text(
        doc,
        "At an optimum, minimization makes t equal to the objective absolute value. The inequalities for u and v ensure that their sum is at least the left side of the original constraint; requiring that sum to be at most five is therefore exact.",
    )


def add_question_6(doc: Document, image_dir: Path, records, solution) -> None:
    doc.add_heading("6. Investment under Taxation", level=1)
    doc.add_heading("6(a) Linear programming formulation", level=2)
    add_text(
        doc,
        "Let the sale variable denote the number of shares of stock i sold now. The net cash coefficient deducts the one percent transaction cost from current proceeds and deducts tax only from a positive capital gain. Since all price data are known, this coefficient is a constant.",
    )
    p = image_dir / "q6_net.png"
    render_formula(
        p,
        [
            r"$g_i=\max\{q_i-p_i,0\},\qquad d_i=0.99q_i-0.30g_i,$",
            r"$d_i=\mathrm{net\ cash\ received\ per\ share\ sold}.$",
        ],
    )
    add_formula(doc, p)
    add_text(doc, "The investor's problem is the following continuous linear program:")
    p = image_dir / "q6_lp.png"
    render_formula(
        p,
        [
            r"$\max_{z}\quad \sum_{i=1}^{n}r_i(s_i-z_i)$",
            r"$\mathrm{s.t.}\quad \sum_{i=1}^{n}d_iz_i\geq K,\qquad 0\leq z_i\leq s_i\quad(i=1,\ldots,n).$",
        ],
    )
    add_formula(doc, p)
    add_text(
        doc,
        "The constant initial portfolio value may be removed from the objective, yielding the equivalent minimization of expected future value sold. The model treats shares as divisible, which is consistent with a linear programming formulation.",
    )

    doc.add_heading("6(b) Numerical solution for K = $9,000", level=2)
    rows = []
    for row, net, ratio in zip(records, solution["net"], solution["ratio"]):
        rows.append(
            [
                row["stock"],
                f"${row['p']:.2f}",
                f"{row['s']:,.0f}",
                f"${row['q']:.2f}",
                f"${row['r']:.2f}",
                f"${net:.3f}",
                f"{ratio:.4f}",
            ]
        )
    make_table(
        doc,
        ["Stock", "Buy price", "Shares", "Current", "Expected", "Net cash/share", "Future loss/$"],
        rows,
        widths=[0.55, 0.78, 0.72, 0.78, 0.78, 1.08, 1.02],
        font_size=7.7,
    )
    add_text(
        doc,
        "The last column is the expected future value forgone per dollar of net cash raised. Its ascending order is S1, S2, S3, S4, then S5. Thus, the first three holdings are sold in full, after which only the amount of S4 needed to reach the target is sold.",
    )
    p = image_dir / "q6_solution_vector.png"
    render_formula(
        p,
        [
            r"$z^*=(1000,\ 1000,\ 1000,\ 118.7618,\ 0)^{\mathsf{T}}\ \mathrm{shares},$",
            r"$\sum_{i=1}^{5}d_iz_i^*=\$9{,}000.00,$",
            r"$\sum_{i=1}^{5}r_i(s_i-z_i^*)=\$15{,}356.79.$",
        ],
        fontsize=15.3,
    )
    add_formula(doc, p)

    result_rows = []
    for row, sold, rem, cash, future in zip(
        records,
        solution["sold"],
        solution["remaining"],
        solution["cash"],
        solution["future_by_stock"],
    ):
        result_rows.append(
            [
                row["stock"],
                f"{sold:,.3f}",
                f"{rem:,.3f}",
                f"${cash:,.2f}",
                f"${future:,.2f}",
            ]
        )
    result_rows.append(
        ["Total", f"{np.sum(solution['sold']):,.3f}", f"{np.sum(solution['remaining']):,.3f}", f"${solution['cash_total']:,.2f}", f"${solution['future_total']:,.2f}"]
    )
    make_table(
        doc,
        ["Stock", "Shares sold", "Shares retained", "Net cash raised", "Expected value retained"],
        result_rows,
        widths=[0.7, 1.1, 1.18, 1.28, 1.52],
        font_size=8.2,
    )
    plot_path = image_dir / "q6_portfolio.png"
    plot_portfolio_solution(plot_path, records, solution)
    add_figure(doc, plot_path, "Figure 3. Optimal sale and retention of shares", width=5.65)
    add_answer_box(doc, "Optimal decision", "Sell all shares of S1, S2, and S3, sell approximately 118.762 shares of S4, and retain all shares of S5. This raises exactly $9,000.00 and leaves an expected one-year portfolio value of $15,356.79.")

    doc.add_heading("Computational code for Question 6(b)", level=2)
    add_text(doc, "The following code is the numerical optimization core used to produce the reported solution.")
    code = '''import csv
import numpy as np
from scipy.optimize import linprog

with open("sources/investment.csv", newline="", encoding="utf-8-sig") as f:
    rows = list(csv.DictReader(f))

p = np.array([float(row["Buying Price"]) for row in rows])
s = np.array([float(row["Current Share"]) for row in rows])
q = np.array([float(row["Current Price"]) for row in rows])
r = np.array([float(row["Expected Future Price"]) for row in rows])

net = 0.99 * q - 0.30 * np.maximum(q - p, 0.0)
result = linprog(
    c=r,
    A_ub=-net.reshape(1, -1),
    b_ub=np.array([-9000.0]),
    bounds=[(0.0, limit) for limit in s],
    method="highs",
)
assert result.success
sold = result.x
cash_raised = net @ sold
future_value = r @ (s - sold)
print(sold, cash_raised, future_value)'''
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GREY)
    set_cell_margins(cell, top=100, start=130, bottom=100, end=130)
    cell.text = ""
    for idx, line in enumerate(code.splitlines()):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.style = "Code"
        p.add_run(line if line else " ")


def build_document(records, solution, image_dir: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = intro.add_run("Solution conventions")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    intro.paragraph_format.space_after = Pt(4)
    add_text(
        doc,
        "All decision variables are continuous unless stated otherwise. Vectors are column vectors, the superscript T denotes transpose, and inequalities between vectors are componentwise. Monetary figures are in dollars.",
    )

    add_question_1(doc, image_dir)
    add_question_2(doc, image_dir)
    add_question_3(doc, image_dir)
    add_question_4(doc, image_dir)
    add_question_5(doc, image_dir)
    add_question_6(doc, image_dir, records, solution)

    core = doc.core_properties
    core.title = "DBA5103 Assignment 1 Solutions"
    core.subject = "Operations Research and Analytics"
    core.author = "Reagan"
    core.keywords = "linear programming, matrix completion, polyhedra, scheduling, taxation"
    core.comments = "Generated from the supplied assignment and investment data."

    # Ensure the document itself contains no em dash characters.
    for paragraph in doc.paragraphs:
        if "\u2014" in paragraph.text:
            raise ValueError("Em dash detected in document text")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "\u2014" in paragraph.text:
                        raise ValueError("Em dash detected in table text")

    doc.save(OUTPUT_FILE)


def validate_output(solution) -> None:
    if not np.isclose(solution["cash_total"], 9000.0, atol=1e-6):
        raise AssertionError("Cash target validation failed")
    if np.any(solution["sold"] < -1e-8) or np.any(solution["sold"] - solution["s"] > 1e-8):
        raise AssertionError("Share-bound validation failed")
    expected = np.array([1000.0, 1000.0, 1000.0, 564.0 / 4.749, 0.0])
    if not np.allclose(solution["sold"], expected, atol=1e-5):
        raise AssertionError("Unexpected investment solution")
    if not OUTPUT_FILE.exists() or OUTPUT_FILE.stat().st_size < 100_000:
        raise AssertionError("DOCX output is missing or unexpectedly small")
    with OUTPUT_FILE.open("rb") as handle:
        if handle.read(2) != b"PK":
            raise AssertionError("DOCX is not a valid ZIP-based Office file")


def main() -> None:
    records = read_portfolio(DATA_FILE)
    solution = solve_investment(records, target=9000.0)
    temp_dir = Path(tempfile.mkdtemp(prefix="dba5103_assignment1_"))
    try:
        build_document(records, solution, temp_dir)
        validate_output(solution)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    print(f"Created: {OUTPUT_FILE}")
    print(f"Solver: {solution['solver_message']}")
    print("Shares sold:", np.round(solution["sold"], 6).tolist())
    print(f"Net cash raised: ${solution['cash_total']:,.2f}")
    print(f"Expected remaining value: ${solution['future_total']:,.2f}")


if __name__ == "__main__":
    main()
