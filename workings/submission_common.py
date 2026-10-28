"""Shared DOCX layout for the DBA5103 assignment submissions.

The styling helpers are reused from the Assignment 1 generator so every
submission has the same look. This module adds the pieces that were edited
into the final Assignment 1 by hand: the full name on the title page, a page
break before every question, the code appendix and the AI declaration.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from assignment1_solution import (
    BLUE,
    MID_GREY,
    NAVY,
    PALE_GREY,
    add_answer_box,
    add_figure,
    add_formula,
    add_text,
    configure_document,
    get_nus_logo,
    make_table as _make_table,
    render_matrix,
    set_cell_margins,
    set_cell_shading,
)
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
STUDENT_NAME = "Reagan (CHAN HIU LOK)"
STUDENT_NO = "A0319086X"
AI_DECLARATION = (
    "I used Claude Opus, OpenAI GPT5.6-Sol to generate ideas and formatting. "
    "I am responsible for the content and quality of the submitted work"
)

__all__ = [
    "BLUE", "MID_GREY", "NAVY", "PALE_GREY", "ROOT",
    "add_answer_box", "add_figure", "add_formula", "add_text", "make_table",
    "render_formula", "render_matrix",
    "new_document", "page_break", "add_code_block", "add_ai_declaration", "save_document",
]


def render_formula(path: Path, lines: str | list[str], fontsize: float = 16.5) -> None:
    """Render equation lines, giving extra height to lines with sum limits or fractions."""
    if isinstance(lines, str):
        lines = [lines]
    heights = [0.72 if ("\\sum" in line or "\\frac" in line or "_{" in line.split("\\quad")[0]) else 0.42
               for line in lines]
    total = sum(heights)
    fig = plt.figure(figsize=(7.1, max(0.44, total)), dpi=300, facecolor="white")
    top = 0.0
    for line, h in zip(lines, heights):
        y = 1 - (top + h / 2) / total
        fig.text(0.5, y, line, ha="center", va="center", fontsize=fontsize, color="#111820")
        top += h
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.07, facecolor="white")
    plt.close(fig)


def make_table(doc: Document, headers, rows, widths=None, font_size=8.5):
    """Styled table that keeps its column widths and never splits across pages."""
    table = _make_table(doc, headers, rows, widths=widths, font_size=font_size)
    tbl = table._tbl
    if widths:
        grid = tbl.tblGrid
        for col, width in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(width * 1440)))
    for idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if idx < len(table.rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    previous = tbl.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        Paragraph(previous, doc).paragraph_format.keep_with_next = True
    return table


def _title_run(p, text, size, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def add_title_page(doc: Document, assignment_no: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(24)
    p.add_run().add_picture(str(get_nus_logo()), width=Inches(2.95))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    _title_run(p, "DBA5103", 14, bold=True, color=BLUE)

    p = doc.add_paragraph("Operations Research and Analytics", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)

    p = doc.add_paragraph(f"Assignment {assignment_no}", style="Title")
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    _title_run(p, "Semester I, Academic Year 2026/2027", 10.5, color=MID_GREY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(3)
    _title_run(p, STUDENT_NAME, 12, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _title_run(p, "MSBA | NUS", 10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # the page break after the title page comes from the first question heading
    _title_run(p, f"Student No: {STUDENT_NO}", 10.5)


def new_document(assignment_no: int) -> Document:
    doc = Document()
    configure_document(doc)
    add_title_page(doc, assignment_no)
    return doc


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GREY)
    set_cell_margins(cell, top=100, start=130, bottom=100, end=130)
    cell.text = ""
    for idx, line in enumerate(code.rstrip().splitlines()):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.style = "Code"
        p.add_run(line if line else " ")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_ai_declaration(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("AI Declaration:")
    r.bold = True
    r.add_break()
    p.add_run(AI_DECLARATION)


def _all_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def save_document(doc: Document, path: Path, title: str, keywords: str) -> None:
    # every question and the appendix start on a new page
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
    for paragraph in _all_paragraphs(doc):
        if "—" in paragraph.text:
            raise ValueError(f"Em dash detected: {paragraph.text[:60]}")
        if "Lecture" in paragraph.text:
            raise ValueError(f"Lecture reference detected: {paragraph.text[:60]}")
    core = doc.core_properties
    core.title = title
    core.subject = "Operations Research and Analytics"
    core.author = "Reagan"
    core.keywords = keywords
    doc.save(path)
    with path.open("rb") as handle:
        if handle.read(2) != b"PK":
            raise AssertionError("DOCX is not a valid ZIP-based Office file")
